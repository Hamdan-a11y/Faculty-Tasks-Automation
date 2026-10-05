import io
from django.test import TestCase, Client
from django.urls import reverse
from openpyxl import Workbook
from .models import StudentRecord


class FinalAcademicInsightsTests(TestCase):
    def setUp(self):
        self.client = Client()

    def create_dummy_excel(self, data, custom_headers=None):
        wb = Workbook()
        ws = wb.active
        
        headers = custom_headers or [
            "Student Name", "Quiz 1", "Quiz 2", "Quiz 3", "Quiz 4",
            "Assign 1", "Assign 2", "Assign 3", "Assign 4",
            "Midterm", "Final", "Total"
        ]
        ws.append(headers)
        
        for row in data:
            ws.append(row)
            
        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        return file_stream

    def test_index_view(self):
        url = reverse("final_academic_insights:index")
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "final_academic_insights/index.html")

    def test_upload_excel_and_flow(self):
        # 3 students: 
        # Student A: total = 4.5*4 (18) + 3*4 (12) + 30 + 40 = 100 -> Pass
        # Student B: total = 2.0*4 (8) + 2.0*4 (8) + 15 + 20 = 51 -> Pass
        # Student C: total = 1.0*4 (4) + 1.0*4 (4) + 10 + 10 = 28 -> Fail
        data = [
            ["Student A", 4.5, 4.5, 4.5, 4.5, 3.0, 3.0, 3.0, 3.0, 30.0, 40.0, 100.0],
            ["Student B", 2.0, 2.0, 2.0, 2.0, 2.0, 2.0, 2.0, 2.0, 15.0, 20.0, 51.0],
            ["Student C", 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 10.0, 10.0, 28.0],
        ]
        
        excel_file = self.create_dummy_excel(data)
        excel_file.name = "marksheet.xlsx"
        
        # Upload POST
        upload_url = reverse("final_academic_insights:upload_excel")
        response = self.client.post(upload_url, {"file": excel_file})
        self.assertEqual(response.status_code, 200)
        res_json = response.json()
        self.assertTrue(res_json["success"])
        
        # Check DB Records
        records = StudentRecord.objects.all()
        self.assertEqual(records.count(), 3)
        
        student_a = StudentRecord.objects.get(student_name="Student A")
        self.assertEqual(student_a.status, "Pass")
        self.assertEqual(student_a.total, 100.0)
        
        student_c = StudentRecord.objects.get(student_name="Student C")
        self.assertEqual(student_c.status, "Fail")
        self.assertEqual(student_c.total, 28.0)
        
        # Test Report View
        report_url = reverse("final_academic_insights:report")
        report_resp = self.client.get(report_url)
        self.assertEqual(report_resp.status_code, 200)
        report_data = report_resp.json()
        
        self.assertEqual(report_data["total_students"], 3)
        self.assertEqual(report_data["passed_students_count"], 2)
        self.assertEqual(report_data["failed_students_count"], 1)
        self.assertEqual(report_data["pass_percentage"], 66.67)
        self.assertEqual(report_data["fail_percentage"], 33.33)
        self.assertEqual(len(report_data["students"]), 3)

        # Verify new enhanced fields
        self.assertAlmostEqual(report_data["average_score"], 59.67, places=1)
        self.assertEqual(report_data["highest_score"], 100.0)
        self.assertIn("grade_distribution", report_data)
        self.assertIn("category_averages", report_data)
        self.assertIn("trend_data", report_data)
        self.assertIn("insights", report_data)
        self.assertTrue(len(report_data["insights"]) >= 4)

        # Verify grade distribution (Student A: 100=A+, Student B: 51=D, Student C: 28=F)
        gd = report_data["grade_distribution"]
        self.assertEqual(gd.get("A+", 0), 1)
        self.assertEqual(gd.get("D", 0), 1)
        self.assertEqual(gd.get("F", 0), 1)

        # Verify category averages exist
        cats = report_data["category_averages"]
        self.assertIn("Quizzes", cats)
        self.assertIn("Assignments", cats)
        self.assertIn("Midterm", cats)
        self.assertIn("Final Exam", cats)

        # Verify trend data has 4 values
        self.assertEqual(len(report_data["trend_data"]), 4)
        
        # Test Chart Data View
        chart_url = reverse("final_academic_insights:chart_data")
        chart_resp = self.client.get(chart_url)
        self.assertEqual(chart_resp.status_code, 200)
        chart_data = chart_resp.json()
        
        self.assertEqual(chart_data["labels"], ["Pass", "Fail"])
        self.assertEqual(chart_data["values"], [2, 1])

    def test_upload_excel_without_total_calculation(self):
        # Student without explicit total column; views must calculate it
        data = [
            ["Student NoTotal", 4.0, 4.0, 4.0, 4.0, 2.0, 2.0, 2.0, 2.0, 15.0, 20.0],
        ]
        headers = [
            "Student Name", "Quiz 1", "Quiz 2", "Quiz 3", "Quiz 4",
            "Assign 1", "Assign 2", "Assign 3", "Assign 4",
            "Midterm", "Final"
        ]
        
        excel_file = self.create_dummy_excel(data, custom_headers=headers)
        excel_file.name = "marksheet_nototal.xlsx"
        
        upload_url = reverse("final_academic_insights:upload_excel")
        response = self.client.post(upload_url, {"file": excel_file})
        self.assertEqual(response.status_code, 200)
        
        record = StudentRecord.objects.get(student_name="Student NoTotal")
        expected_total = 4.0*4 + 2.0*4 + 15.0 + 20.0  # 16 + 8 + 15 + 20 = 59.0
        self.assertEqual(record.total, expected_total)
        self.assertEqual(record.status, "Pass")

    def test_invalid_file_upload(self):
        # Non-excel file
        file_stream = io.BytesIO(b"hello world")
        file_stream.name = "test.txt"
        
        upload_url = reverse("final_academic_insights:upload_excel")
        response = self.client.post(upload_url, {"file": file_stream})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.json()["success"])
        
    def test_missing_required_columns(self):
        # Excel missing Quiz 1 column
        data = [
            ["Student X", 4.0, 3.0, 20.0, 30.0],
        ]
        headers = ["Student Name", "Quiz 2", "Assignment 1", "Midterm", "Final"]
        excel_file = self.create_dummy_excel(data, custom_headers=headers)
        excel_file.name = "missing_cols.xlsx"
        
        upload_url = reverse("final_academic_insights:upload_excel")
        response = self.client.post(upload_url, {"file": excel_file})
        self.assertEqual(response.status_code, 400)
        self.assertIn("Too many missing columns", response.json()["error"])

    def test_download_excel_report(self):
        """Verify Excel download endpoint returns valid .xlsx file."""
        # Upload data first
        data = [
            ["Student A", 4.5, 4.5, 4.5, 4.5, 3.0, 3.0, 3.0, 3.0, 30.0, 40.0, 100.0],
            ["Student B", 2.0, 2.0, 2.0, 2.0, 2.0, 2.0, 2.0, 2.0, 15.0, 20.0, 51.0],
            ["Student C", 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 10.0, 10.0, 28.0],
        ]
        excel_file = self.create_dummy_excel(data)
        excel_file.name = "marksheet.xlsx"
        upload_url = reverse("final_academic_insights:upload_excel")
        self.client.post(upload_url, {"file": excel_file})

        # Download Excel report
        download_url = reverse("final_academic_insights:download_excel")
        response = self.client.get(download_url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn("final_academic_insights_report.xlsx", response["Content-Disposition"])

        # Verify the file is valid by opening with openpyxl
        import openpyxl
        file_stream = io.BytesIO(response.content)
        wb = openpyxl.load_workbook(file_stream)
        sheet_names = wb.sheetnames
        self.assertIn("Academic Overview", sheet_names)
        self.assertIn("Grade Distribution", sheet_names)
        self.assertIn("Assessment Breakdown", sheet_names)
        self.assertIn("Academic Insights", sheet_names)
        self.assertIn("Student Records", sheet_names)

    def test_download_excel_empty(self):
        """Verify download returns 400 when no data exists."""
        download_url = reverse("final_academic_insights:download_excel")
        response = self.client.get(download_url)
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.json()["success"])
