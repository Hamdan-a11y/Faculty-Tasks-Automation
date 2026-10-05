"""Midterm Academic Evaluator views and analytics."""

import os
import csv
import json
import re
import uuid
import tempfile
import statistics
from io import BytesIO
from datetime import datetime

from django.conf import settings
from django.http import JsonResponse, HttpResponseBadRequest, HttpResponse
from django.shortcuts import render, redirect
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.db import DatabaseError, transaction

from auth_app.models import (
    FacultyUser,
    MidtermEvaluationRun,
    MidtermEvaluationStudent,
)
from .helpers import (
    _get_session_user,
    _normalize_header_text,
    _is_registration_header,
    _is_student_name_header,
    SESSION_KEY,
)

MIDTERM_GRADE_ORDER = ["A", "A-", "B+", "B", "B-", "C+", "C", "C-", "D", "F"]


MIDTERM_PERFORMANCE_ORDER = ["Excellent", "Good", "Satisfactory", "At Risk", "Critical"]


MIDTERM_GRADE_SCALE = [
    (90.0, "A", 4.0),
    (85.0, "A-", 3.7),
    (80.0, "B+", 3.3),
    (75.0, "B", 3.0),
    (70.0, "B-", 2.7),
    (65.0, "C+", 2.3),
    (60.0, "C", 2.0),
    (55.0, "C-", 1.7),
    (50.0, "D", 1.0),
    (0.0, "F", 0.0),
]


def _coerce_midterm_numeric(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return float(int(value))
    if isinstance(value, (int, float)):
        return round(float(value), 2)

    text = str(value).strip()
    if not text:
        return None

    if "/" in text:
        text = text.split("/")[0].strip()

    text = text.replace(",", "")
    try:
        return round(float(text), 2)
    except ValueError:
        numeric_match = re.search(r"-?\d+(?:\.\d+)?", text)
        if not numeric_match:
            return None
        return round(float(numeric_match.group(0)), 2)


def _is_midterm_marks_header(header_key):
    if not header_key:
        return False
    aliases = {
        "midterm",
        "midtermpaper",
        "midtermexam",
        "midtermmarks",
        "midtermscore",
        "midterm30marks",
        "mid",
    }
    if header_key in aliases:
        return True
    if "mid" in header_key and ("term" in header_key or "paper" in header_key or "exam" in header_key):
        return True
    return False


def _extract_marks_cap(text_value):
    text = str(text_value or "")
    cap_match = re.search(r"(\d+(?:\.\d+)?)\s*marks?", text, flags=re.IGNORECASE)
    if not cap_match:
        return None
    try:
        return float(cap_match.group(1))
    except ValueError:
        return None


def _parse_midterm_students_from_file(file_path):
    try:
        from openpyxl import load_workbook
    except Exception as import_error:
        raise RuntimeError(f"Excel parser dependency is missing: {import_error}") from import_error

    workbook = load_workbook(filename=file_path, data_only=True, read_only=True)
    sheet_candidates = []
    try:
        for worksheet in workbook.worksheets:
            rows = []
            for row_number, row_values in enumerate(worksheet.iter_rows(values_only=True), start=1):
                rows.append({"row_number": row_number, "values": list(row_values)})
            sheet_candidates.append({"sheet_name": worksheet.title, "rows": rows})
    finally:
        workbook.close()

    best_match = None
    for candidate_sheet in sheet_candidates:
        rows = candidate_sheet.get("rows") or []
        if not rows:
            continue

        header_row_number = None
        registration_col = None
        name_col = None
        midterm_col = None
        header_values = []

        for candidate in rows[:120]:
            values = candidate.get("values", [])
            normalized_headers = [_normalize_header_text(value) for value in values]
            reg_idx = next((i for i, key in enumerate(normalized_headers) if _is_registration_header(key)), None)
            name_idx = next((i for i, key in enumerate(normalized_headers) if _is_student_name_header(key)), None)
            mid_idx = next((i for i, key in enumerate(normalized_headers) if _is_midterm_marks_header(key)), None)

            if reg_idx is not None and name_idx is not None and mid_idx is not None:
                header_row_number = candidate.get("row_number")
                registration_col = reg_idx
                name_col = name_idx
                midterm_col = mid_idx
                header_values = values
                break

        if header_row_number is None:
            continue

        candidate_student_count = 0
        candidate_numeric_marks = 0
        for row_obj in rows:
            row_number = row_obj.get("row_number") or 0
            if row_number <= header_row_number:
                continue

            values = row_obj.get("values", [])
            registration_raw = values[registration_col] if registration_col < len(values) else None
            name_raw = values[name_col] if name_col < len(values) else None
            mark_raw = values[midterm_col] if midterm_col < len(values) else None

            if not str(registration_raw or "").strip() and not str(name_raw or "").strip():
                continue

            candidate_student_count += 1
            if _coerce_midterm_numeric(mark_raw) is not None:
                candidate_numeric_marks += 1

        score = (candidate_numeric_marks, candidate_student_count)
        if best_match is None or score > best_match["score"]:
            best_match = {
                "sheet_name": candidate_sheet.get("sheet_name"),
                "rows": rows,
                "header_row": header_row_number,
                "registration_col": registration_col,
                "name_col": name_col,
                "midterm_col": midterm_col,
                "header_values": header_values,
                "score": score,
            }

    if best_match is None:
        raise ValueError(
            "Unable to detect Student Name, Registration Number, and Midterm Marks columns in uploaded sheet."
        )

    students = []
    numeric_marks = []
    for row_obj in best_match["rows"]:
        row_number = row_obj.get("row_number") or 0
        if row_number <= best_match["header_row"]:
            continue

        values = row_obj.get("values", [])
        registration_raw = values[best_match["registration_col"]] if best_match["registration_col"] < len(values) else None
        name_raw = values[best_match["name_col"]] if best_match["name_col"] < len(values) else None
        mark_raw = values[best_match["midterm_col"]] if best_match["midterm_col"] < len(values) else None

        registration_no = str(registration_raw).strip() if registration_raw is not None else ""
        student_name = str(name_raw).strip() if name_raw is not None else ""
        midterm_mark = _coerce_midterm_numeric(mark_raw)

        if not registration_no and not student_name:
            continue

        if midterm_mark is not None:
            numeric_marks.append(midterm_mark)

        students.append(
            {
                "row_number": row_number,
                "registration_no": registration_no,
                "student_name": student_name,
                "midterm_marks": midterm_mark,
            }
        )

    if not students:
        raise ValueError("No student records were found after header detection.")

    header_values = best_match.get("header_values") or []
    midterm_col = best_match["midterm_col"]
    header_for_midterm = header_values[midterm_col] if midterm_col < len(header_values) else ""
    max_marks = _extract_marks_cap(header_for_midterm) or 30.0

    if numeric_marks:
        observed_max = max(numeric_marks)
        if observed_max > max_marks:
            max_marks = observed_max

    return {
        "sheet_name": best_match["sheet_name"],
        "header_row": best_match["header_row"],
        "registration_column": best_match["registration_col"] + 1,
        "name_column": best_match["name_col"] + 1,
        "midterm_column": best_match["midterm_col"] + 1,
        "max_marks": round(float(max_marks), 2),
        "students": students,
    }


def _grade_for_percentage(percent):
    for min_percent, grade, gpa in MIDTERM_GRADE_SCALE:
        if percent >= min_percent:
            return grade, gpa
    return "F", 0.0


def _performance_band_for_score(mark_value, percent):
    if mark_value is None:
        return "At Risk"
    if mark_value < 10:
        return "Critical"
    if percent >= 85:
        return "Excellent"
    if percent >= 70:
        return "Good"
    if percent >= 50:
        return "Satisfactory"
    return "At Risk"


def _risk_status_for_score(mark_value, percent):
    if mark_value is None:
        return "At Risk - Mark Missing"
    if mark_value < 10:
        return "Critical - High Risk of Failing"
    if percent < 50:
        return "At Risk"
    return "Normal"


def _evaluate_midterm_run(run):
    students = list(run.students.all().order_by("row_number", "id"))
    if not students:
        raise ValueError("No extracted student records found. Please extract the sheet first.")

    max_marks = float(run.max_marks or 30.0)
    if max_marks <= 0:
        max_marks = 30.0

    pass_threshold = max_marks * 0.5
    pass_count = 0
    fail_count = 0
    critical_count = 0

    grade_distribution = {grade: 0 for grade in MIDTERM_GRADE_ORDER}
    performance_distribution = {band: 0 for band in MIDTERM_PERFORMANCE_ORDER}

    numeric_marks = []
    highest_entry = None
    lowest_entry = None
    updates = []

    for student in students:
        mark_value = student.midterm_marks
        if mark_value is not None:
            mark_value = round(float(mark_value), 2)
            student.midterm_marks = mark_value

        percent = round((mark_value / max_marks) * 100, 2) if mark_value is not None else 0.0
        grade_letter, gpa_value = _grade_for_percentage(percent)
        performance_band = _performance_band_for_score(mark_value, percent)
        risk_status = _risk_status_for_score(mark_value, percent)
        is_pass = mark_value is not None and mark_value >= pass_threshold

        if is_pass:
            pass_count += 1
        else:
            fail_count += 1

        if performance_band == "Critical":
            critical_count += 1

        grade_distribution[grade_letter] = grade_distribution.get(grade_letter, 0) + 1
        performance_distribution[performance_band] = performance_distribution.get(performance_band, 0) + 1

        student.grade_letter = grade_letter
        student.gpa = float(gpa_value)
        student.performance_band = performance_band
        student.risk_status = risk_status
        student.is_pass = is_pass
        updates.append(student)

        if mark_value is not None:
            numeric_marks.append(mark_value)
            if highest_entry is None or mark_value > highest_entry["marks"]:
                highest_entry = {
                    "marks": mark_value,
                    "student_name": student.student_name,
                    "registration_no": student.registration_no,
                }
            if lowest_entry is None or mark_value < lowest_entry["marks"]:
                lowest_entry = {
                    "marks": mark_value,
                    "student_name": student.student_name,
                    "registration_no": student.registration_no,
                }

    MidtermEvaluationStudent.objects.bulk_update(
        updates,
        ["midterm_marks", "grade_letter", "gpa", "performance_band", "risk_status", "is_pass"],
        batch_size=500,
    )

    average_marks = round(sum(numeric_marks) / len(numeric_marks), 2) if numeric_marks else None
    median_marks = round(float(statistics.median(numeric_marks)), 2) if numeric_marks else None
    total_students = len(students)
    pass_percentage = round((pass_count / total_students) * 100, 2) if total_students else 0.0

    run.total_students = total_students
    run.pass_count = pass_count
    run.fail_count = fail_count
    run.pass_percentage = pass_percentage
    run.highest_marks = highest_entry["marks"] if highest_entry else None
    run.highest_student_name = highest_entry["student_name"] if highest_entry else ""
    run.lowest_marks = lowest_entry["marks"] if lowest_entry else None
    run.lowest_student_name = lowest_entry["student_name"] if lowest_entry else ""
    run.average_marks = average_marks
    run.median_marks = median_marks
    run.grade_distribution = grade_distribution
    run.performance_distribution = performance_distribution
    run.status = "evaluated"
    run.error_message = ""
    run.save(
        update_fields=[
            "total_students",
            "pass_count",
            "fail_count",
            "pass_percentage",
            "highest_marks",
            "highest_student_name",
            "lowest_marks",
            "lowest_student_name",
            "average_marks",
            "median_marks",
            "grade_distribution",
            "performance_distribution",
            "status",
            "error_message",
            "updated_at",
        ]
    )

    return {
        "max_marks": round(max_marks, 2),
        "highest": highest_entry,
        "lowest": lowest_entry,
        "average_marks": average_marks,
        "median_marks": median_marks,
        "pass_count": pass_count,
        "fail_count": fail_count,
        "pass_percentage": pass_percentage,
        "fail_percentage": round(100.0 - pass_percentage, 2),
        "grade_distribution": grade_distribution,
        "performance_distribution": performance_distribution,
        "critical_count": critical_count,
    }


def _serialize_midterm_students(students):
    payload = []
    for student in students:
        payload.append(
            {
                "row_number": student.row_number,
                "registration_no": student.registration_no,
                "student_name": student.student_name,
                "midterm_marks": student.midterm_marks,
                "grade": student.grade_letter,
                "gpa": student.gpa,
                "performance_band": student.performance_band,
                "risk_status": student.risk_status,
                "is_pass": student.is_pass,
            }
        )
    return payload


def _resolve_midterm_run(request, run_id=None):
    user = _get_session_user(request)
    queryset = MidtermEvaluationRun.objects.all()
    if user:
        queryset = queryset.filter(user=user)

    if run_id:
        return user, queryset.filter(id=run_id).first()

    session_run_id = request.session.get("midterm_run_id")
    if session_run_id:
        run = queryset.filter(id=session_run_id).first()
        if run:
            return user, run

    return user, queryset.order_by("-created_at").first()


def midterm_academic_evaluator_page(request):
    user_id = request.session.get(SESSION_KEY)
    user_name = None
    if user_id:
        try:
            user = FacultyUser.objects.get(id=user_id)
            user_name = user.name or user.email
        except FacultyUser.DoesNotExist:
            pass
    context = {"user_name": user_name}
    return render(request, "auth/midterm_academic_evaluator.html", context)


@csrf_exempt
def midterm_upload_excel(request):
    if request.method != "POST":
        return JsonResponse({"error": "Method not allowed"}, status=405)

    uploaded = request.FILES.get("file")
    if not uploaded:
        return JsonResponse({"error": "No file provided."}, status=400)

    _, ext = os.path.splitext(uploaded.name or "")
    ext = ext.lower()
    if ext != ".xlsx":
        return JsonResponse(
            {"error": "Unsupported file type. Please upload a .xlsx file."},
            status=400,
        )

    target_dir = os.path.join(tempfile.gettempdir(), "midterm_excel_uploads")
    os.makedirs(target_dir, exist_ok=True)

    temp_name = f"midterm_eval_{uuid.uuid4().hex}{ext}"
    temp_path = os.path.join(target_dir, temp_name)
    user = _get_session_user(request)

    try:
        with open(temp_path, "wb") as tmp:
            for chunk in uploaded.chunks():
                tmp.write(chunk)

        old_path = request.session.get("midterm_uploaded_excel_path")
        if old_path and os.path.exists(old_path):
            try:
                os.remove(old_path)
            except Exception:
                pass

        run = MidtermEvaluationRun.objects.create(
            user=user,
            source_file_name=uploaded.name,
            source_file_path=temp_path,
            source_file_type=ext,
            status="uploaded",
        )

        request.session["midterm_uploaded_excel_path"] = temp_path
        request.session["midterm_uploaded_excel_name"] = uploaded.name
        request.session["midterm_run_id"] = run.id
        request.session.modified = True

        return JsonResponse(
            {
                "status": "uploaded",
                "run_id": run.id,
                "file_name": uploaded.name,
                "message": "Midterm sheet uploaded successfully.",
            }
        )
    except Exception as exc:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass
        return JsonResponse(
            {
                "error": "Unable to upload the file. Please try again.",
                "detail": str(exc),
            },
            status=500,
        )


@csrf_exempt
def midterm_extract_data(request):
    if request.method not in ["GET", "POST"]:
        return JsonResponse({"error": "Method not allowed"}, status=405)

    run_id = request.GET.get("run_id")
    if request.method == "POST":
        try:
            payload = json.loads(request.body.decode("utf-8") or "{}")
            run_id = run_id or payload.get("run_id")
        except json.JSONDecodeError:
            return JsonResponse({"error": "Invalid JSON payload."}, status=400)

    user, run = _resolve_midterm_run(request, run_id)
    file_path = request.session.get("midterm_uploaded_excel_path")
    file_name = request.session.get("midterm_uploaded_excel_name")

    if run and run.source_file_path:
        file_path = run.source_file_path
        file_name = run.source_file_name or file_name

    if not file_path:
        return JsonResponse({"error": "No midterm file uploaded yet. Please upload first."}, status=400)

    if not os.path.exists(file_path):
        return JsonResponse({"error": "Uploaded file no longer exists. Please upload again."}, status=400)

    try:
        extracted = _parse_midterm_students_from_file(file_path)
        student_rows = extracted.get("students", [])
        if not student_rows:
            raise ValueError("No student data found in the uploaded sheet.")

        with transaction.atomic():
            if run is None:
                run = MidtermEvaluationRun.objects.create(
                    user=user,
                    source_file_name=file_name or "",
                    source_file_path=file_path,
                    source_file_type=".xlsx",
                    status="uploaded",
                )

            run.students.all().delete()
            MidtermEvaluationStudent.objects.bulk_create(
                [
                    MidtermEvaluationStudent(
                        run=run,
                        row_number=row.get("row_number") or 0,
                        registration_no=str(row.get("registration_no") or "").strip(),
                        student_name=str(row.get("student_name") or "").strip(),
                        midterm_marks=row.get("midterm_marks"),
                    )
                    for row in student_rows
                ],
                batch_size=500,
            )

            run.source_file_name = file_name or run.source_file_name
            run.source_file_path = file_path
            run.source_file_type = ".xlsx"
            run.source_sheet_name = extracted.get("sheet_name") or ""
            run.max_marks = extracted.get("max_marks") or run.max_marks or 30.0
            run.status = "extracted"
            run.error_message = ""
            run.save(
                update_fields=[
                    "source_file_name",
                    "source_file_path",
                    "source_file_type",
                    "source_sheet_name",
                    "max_marks",
                    "status",
                    "error_message",
                    "updated_at",
                ]
            )

        summary = _evaluate_midterm_run(run)
        students_payload = _serialize_midterm_students(run.students.all().order_by("row_number", "id"))

        request.session["midterm_run_id"] = run.id
        request.session.modified = True

        return JsonResponse(
            {
                "status": "success",
                "message": "Midterm data extracted and evaluated successfully.",
                "run_id": run.id,
                "file_name": run.source_file_name,
                "sheet_name": run.source_sheet_name,
                "header_row": extracted.get("header_row"),
                "registration_column": extracted.get("registration_column"),
                "name_column": extracted.get("name_column"),
                "midterm_column": extracted.get("midterm_column"),
                "preview_students": students_payload,
                "summary": summary,
            }
        )
    except ValueError as exc:
        if run:
            run.status = "failed"
            run.error_message = str(exc)
            run.save(update_fields=["status", "error_message", "updated_at"])
        return JsonResponse({"error": str(exc)}, status=400)
    except Exception as exc:
        if run:
            run.status = "failed"
            run.error_message = str(exc)
            run.save(update_fields=["status", "error_message", "updated_at"])
        return JsonResponse(
            {
                "error": "Failed to extract and evaluate midterm data. Please verify your file format and retry.",
                "detail": str(exc),
            },
            status=500,
        )


@csrf_exempt
def midterm_evaluate_data(request):
    if request.method != "POST":
        return JsonResponse({"error": "Method not allowed"}, status=405)

    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        return JsonResponse({"error": "Invalid JSON payload."}, status=400)

    run_id = payload.get("run_id")
    _, run = _resolve_midterm_run(request, run_id)
    if not run:
        return JsonResponse({"error": "No extracted midterm run found. Please upload and extract first."}, status=400)

    try:
        summary = _evaluate_midterm_run(run)
        students_payload = _serialize_midterm_students(run.students.all().order_by("row_number", "id"))
        request.session["midterm_run_id"] = run.id
        request.session.modified = True
        return JsonResponse(
            {
                "status": "success",
                "message": "Midterm evaluation recalculated successfully.",
                "run_id": run.id,
                "summary": summary,
                "preview_students": students_payload,
            }
        )
    except ValueError as exc:
        return JsonResponse({"error": str(exc)}, status=400)
    except Exception as exc:
        return JsonResponse(
            {
                "error": "Failed to evaluate student data.",
                "detail": str(exc),
            },
            status=500,
        )


@csrf_exempt
def midterm_generate_report(request):
    if request.method not in ["GET", "POST"]:
        return JsonResponse({"error": "Method not allowed"}, status=405)

    run_id = request.GET.get("run_id")
    if request.method == "POST":
        try:
            payload = json.loads(request.body.decode("utf-8") or "{}")
            run_id = run_id or payload.get("run_id")
        except json.JSONDecodeError:
            return JsonResponse({"error": "Invalid JSON payload."}, status=400)

    _, run = _resolve_midterm_run(request, run_id)
    if not run:
        return JsonResponse({"error": "No evaluation run found. Please upload and extract first."}, status=400)

    students = list(run.students.all().order_by("row_number", "id"))
    if not students:
        return JsonResponse({"error": "No student records available for report generation."}, status=400)

    try:
        if run.status != "evaluated":
            _evaluate_midterm_run(run)
            students = list(run.students.all().order_by("row_number", "id"))

        from openpyxl import Workbook
        from openpyxl.chart import BarChart, PieChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"

        header_fill = PatternFill(start_color="DFF3FB", end_color="DFF3FB", fill_type="solid")
        title_font = Font(bold=True, size=13, color="0D3B57")
        header_font = Font(bold=True, color="0D3B57")
        critical_fill = PatternFill(start_color="FDE2E1", end_color="FDE2E1", fill_type="solid")

        summary_sheet["A1"] = "Midterm Academic Evaluation Summary"
        summary_sheet["A1"].font = title_font
        summary_sheet.merge_cells("A1:D1")

        highest_label = "N/A"
        if run.highest_marks is not None:
            highest_label = f"{run.highest_student_name} ({run.highest_marks})"

        lowest_label = "N/A"
        if run.lowest_marks is not None:
            lowest_label = f"{run.lowest_student_name} ({run.lowest_marks})"

        summary_rows = [
            ("File Name", run.source_file_name or "N/A"),
            ("Sheet", run.source_sheet_name or "N/A"),
            ("Maximum Marks", run.max_marks),
            ("Total Students", run.total_students),
            ("Highest Marks", highest_label),
            ("Lowest Marks", lowest_label),
            ("Average Marks", run.average_marks if run.average_marks is not None else "N/A"),
            ("Median Marks", run.median_marks if run.median_marks is not None else "N/A"),
            ("Pass Count", run.pass_count),
            ("Fail Count", run.fail_count),
            ("Pass Percentage", f"{run.pass_percentage}%"),
        ]

        row_cursor = 3
        for key, value in summary_rows:
            summary_sheet[f"A{row_cursor}"] = key
            summary_sheet[f"B{row_cursor}"] = value
            summary_sheet[f"A{row_cursor}"].font = header_font
            row_cursor += 1

        row_cursor += 1
        summary_sheet[f"A{row_cursor}"] = "Grade Distribution"
        summary_sheet[f"A{row_cursor}"].font = title_font
        row_cursor += 1
        grade_header_row = row_cursor
        summary_sheet[f"A{row_cursor}"] = "Grade"
        summary_sheet[f"B{row_cursor}"] = "Count"
        summary_sheet[f"A{row_cursor}"].font = header_font
        summary_sheet[f"B{row_cursor}"].font = header_font
        summary_sheet[f"A{row_cursor}"].fill = header_fill
        summary_sheet[f"B{row_cursor}"].fill = header_fill

        grade_data_start_row = row_cursor + 1
        for grade in MIDTERM_GRADE_ORDER:
            row_cursor += 1
            summary_sheet[f"A{row_cursor}"] = grade
            summary_sheet[f"B{row_cursor}"] = int((run.grade_distribution or {}).get(grade, 0))
        grade_data_end_row = row_cursor

        row_cursor += 2
        summary_sheet[f"A{row_cursor}"] = "Performance Bands"
        summary_sheet[f"A{row_cursor}"].font = title_font
        row_cursor += 1
        band_header_row = row_cursor
        summary_sheet[f"A{row_cursor}"] = "Band"
        summary_sheet[f"B{row_cursor}"] = "Count"
        summary_sheet[f"A{row_cursor}"].font = header_font
        summary_sheet[f"B{row_cursor}"].font = header_font
        summary_sheet[f"A{row_cursor}"].fill = header_fill
        summary_sheet[f"B{row_cursor}"].fill = header_fill

        band_data_start_row = row_cursor + 1
        for band in MIDTERM_PERFORMANCE_ORDER:
            row_cursor += 1
            summary_sheet[f"A{row_cursor}"] = band
            summary_sheet[f"B{row_cursor}"] = int((run.performance_distribution or {}).get(band, 0))
        band_data_end_row = row_cursor

        chart_source_header_row = row_cursor + 2
        summary_sheet[f"H{chart_source_header_row}"] = "Result"
        summary_sheet[f"I{chart_source_header_row}"] = "Count"
        summary_sheet[f"H{chart_source_header_row}"].font = header_font
        summary_sheet[f"I{chart_source_header_row}"].font = header_font
        summary_sheet[f"H{chart_source_header_row}"].fill = header_fill
        summary_sheet[f"I{chart_source_header_row}"].fill = header_fill
        summary_sheet[f"H{chart_source_header_row + 1}"] = "Pass"
        summary_sheet[f"I{chart_source_header_row + 1}"] = int(run.pass_count or 0)
        summary_sheet[f"H{chart_source_header_row + 2}"] = "Fail"
        summary_sheet[f"I{chart_source_header_row + 2}"] = int(run.fail_count or 0)

        grade_chart = BarChart()
        grade_chart.type = "col"
        grade_chart.style = 10
        grade_chart.title = "Grade Distribution"
        grade_chart.y_axis.title = "Students"
        grade_chart.x_axis.title = "Grade"
        grade_chart.add_data(
            Reference(summary_sheet, min_col=2, min_row=grade_header_row, max_row=grade_data_end_row),
            titles_from_data=True,
        )
        grade_chart.set_categories(
            Reference(summary_sheet, min_col=1, min_row=grade_data_start_row, max_row=grade_data_end_row)
        )
        grade_chart.height = 7
        grade_chart.width = 12
        summary_sheet.add_chart(grade_chart, "D3")

        performance_chart = BarChart()
        performance_chart.type = "col"
        performance_chart.style = 10
        performance_chart.title = "Performance Band Distribution"
        performance_chart.y_axis.title = "Students"
        performance_chart.x_axis.title = "Band"
        performance_chart.add_data(
            Reference(summary_sheet, min_col=2, min_row=band_header_row, max_row=band_data_end_row),
            titles_from_data=True,
        )
        performance_chart.set_categories(
            Reference(summary_sheet, min_col=1, min_row=band_data_start_row, max_row=band_data_end_row)
        )
        performance_chart.height = 7
        performance_chart.width = 12
        summary_sheet.add_chart(performance_chart, "D20")

        pass_fail_chart = PieChart()
        pass_fail_chart.title = "Pass vs Fail"
        pass_fail_chart.add_data(
            Reference(summary_sheet, min_col=9, min_row=chart_source_header_row, max_row=chart_source_header_row + 2),
            titles_from_data=True,
        )
        pass_fail_chart.set_categories(
            Reference(summary_sheet, min_col=8, min_row=chart_source_header_row + 1, max_row=chart_source_header_row + 2)
        )
        pass_fail_chart.height = 6
        pass_fail_chart.width = 9
        summary_sheet.add_chart(pass_fail_chart, "D37")

        detail_sheet = workbook.create_sheet(title="Student Detail")
        detail_headers = [
            "Registration No",
            "Student Name",
            "Midterm Marks",
            "Grade",
            "GPA",
            "Performance Band",
            "Risk Status",
            "Pass/Fail",
        ]
        detail_sheet.append(detail_headers)
        for col_idx in range(1, len(detail_headers) + 1):
            cell = detail_sheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for student in students:
            detail_sheet.append(
                [
                    student.registration_no,
                    student.student_name,
                    student.midterm_marks,
                    student.grade_letter,
                    student.gpa,
                    student.performance_band,
                    student.risk_status,
                    "Pass" if student.is_pass else "Fail",
                ]
            )

            if (student.risk_status or "").lower().startswith("critical"):
                row_idx = detail_sheet.max_row
                for col_idx in range(1, len(detail_headers) + 1):
                    detail_sheet.cell(row=row_idx, column=col_idx).fill = critical_fill

        risk_sheet = workbook.create_sheet(title="Risk Report")
        risk_headers = [
            "Registration No",
            "Student Name",
            "Midterm Marks",
            "Grade",
            "Performance Band",
            "Risk Status",
        ]
        risk_sheet.append(risk_headers)
        for col_idx in range(1, len(risk_headers) + 1):
            cell = risk_sheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        risk_students = [
            student
            for student in students
            if student.performance_band in {"At Risk", "Critical"}
            or (student.risk_status or "").lower().startswith("critical")
        ]

        if not risk_students:
            risk_sheet.append(["-", "No at-risk or critical students found", "-", "-", "-", "-"])
        else:
            for student in risk_students:
                risk_sheet.append(
                    [
                        student.registration_no,
                        student.student_name,
                        student.midterm_marks,
                        student.grade_letter,
                        student.performance_band,
                        student.risk_status,
                    ]
                )
                if (student.risk_status or "").lower().startswith("critical"):
                    row_idx = risk_sheet.max_row
                    for col_idx in range(1, len(risk_headers) + 1):
                        risk_sheet.cell(row=row_idx, column=col_idx).fill = critical_fill

        for sheet in [summary_sheet, detail_sheet, risk_sheet]:
            for column_cells in sheet.columns:
                max_length = 0
                column_index = column_cells[0].column
                for cell in column_cells:
                    value = cell.value
                    if value is None:
                        continue
                    max_length = max(max_length, len(str(value)))
                sheet.column_dimensions[get_column_letter(column_index)].width = min(max_length + 3, 48)

        timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
        report_file_name = f"midterm_evaluation_report_{run.id}_{timestamp}.xlsx"

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        run.report_file_name = report_file_name
        run.report_generated_at = timezone.now()
        run.status = "reported"
        run.save(update_fields=["report_file_name", "report_generated_at", "status", "updated_at"])

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{report_file_name}"'
        return response
    except ValueError as exc:
        return JsonResponse({"error": str(exc)}, status=400)
    except Exception as exc:
        return JsonResponse(
            {
                "error": "Failed to generate the evaluation report.",
                "detail": str(exc),
            },
            status=500,
        )

