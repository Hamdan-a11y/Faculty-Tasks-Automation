"""ZABDESK marks recap sheet entry and Excel data processing."""

import os
import csv
import re
import time
import json
import uuid
import tempfile
import subprocess
from datetime import datetime

from django.conf import settings
from django.http import JsonResponse, HttpResponseBadRequest
from django.shortcuts import render, redirect
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.db import DatabaseError, transaction
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service as ChromeService
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    TimeoutException,
    NoSuchElementException,
    WebDriverException,
    InvalidSessionIdException,
)

from auth_app.models import (
    FacultyUser,
    ZabdeskMarkEntryRun,
)
from .helpers import (
    _get_session_user,
    _is_remote_debugger_running,
    get_driver_path,
    _find_chrome_executable,
    _normalize_header_text,
    _is_registration_header,
    _is_student_name_header,
    CHROME_PROFILE_PATH,
    LOGIN_URL,
    SESSION_KEY,
)

def zabdesk_marks_bridge_page(request):
    user_id = request.session.get(SESSION_KEY)
    user_name = None
    if user_id:
        try:
            user = FacultyUser.objects.get(id=user_id)
            user_name = user.name or user.email
        except FacultyUser.DoesNotExist:
            pass
    context = {"user_name": user_name}
    return render(request, "auth/zabdesk_marks_bridge.html", context)


def zabdesk_auto_login(request):
    if request.method not in ["GET", "POST"]:
        return JsonResponse({"error": "Method not allowed"}, status=405)
    try:
        portal_root = "https://hamdan-a11y.github.io/ZABDESK-Dummy-Portal/"

        def wait_for_url(tokens, timeout=12):
            return WebDriverWait(driver, timeout).until(
                lambda d: any(token in d.current_url.lower() for token in tokens)
            )

        def click_when_clickable(xpath, label, timeout=10):
            try:
                element = WebDriverWait(driver, timeout).until(
                    EC.element_to_be_clickable((By.XPATH, xpath))
                )
                driver.execute_script("arguments[0].click();", element)
                return True
            except TimeoutException:
                print(f"Zabdesk Auto: could not click {label}")
                return False

        options = Options()
        options.add_experimental_option("debuggerAddress", "127.0.0.1:9222")
        chrome_path = _find_chrome_executable()

        if not chrome_path:
            return JsonResponse(
                {
                    "status": "error",
                    "message": "Chrome executable not found. Please install Google Chrome.",
                },
                status=500,
            )

        if not _is_remote_debugger_running():
            subprocess.Popen([
                chrome_path, "--remote-debugging-port=9222", "--remote-allow-origins=*",
                f"--user-data-dir={CHROME_PROFILE_PATH}", "--profile-directory=Default",
                "--no-first-run", "--no-default-browser-check"
            ], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            for _ in range(60):
                if _is_remote_debugger_running(): break
                time.sleep(0.1)

            if not _is_remote_debugger_running():
                return JsonResponse(
                    {
                        "status": "error",
                        "message": "Chrome remote debugging session could not be started.",
                    },
                    status=500,
                )

        service = ChromeService(executable_path=get_driver_path())
        driver = webdriver.Chrome(service=service, options=options)

        # Reuse an existing ZabDesk tab when possible; avoid chrome:// pages.
        target_handle = None
        for handle in driver.window_handles:
            driver.switch_to.window(handle)
            if "zabdesk-dummy-portal" in driver.current_url.lower():
                target_handle = handle
                break

        if not target_handle:
            for handle in driver.window_handles:
                driver.switch_to.window(handle)
                if not driver.current_url.lower().startswith("chrome://"):
                    target_handle = handle
                    break

        if not target_handle:
            driver.switch_to.new_window("tab")
        else:
            driver.switch_to.window(target_handle)

        driver.get(portal_root)
        WebDriverWait(driver, 15).until(
            lambda d: d.execute_script("return document.readyState") == "complete"
        )

        current_url = driver.current_url.lower()
        already_authenticated = any(
            token in current_url
            for token in ["roles.html", "faculty-courses.html", "course-hub.html", "recap-sheet.html"]
        )

        if not already_authenticated:
            username_input = None
            password_input = None

            username_candidates = [
                "//main//input[@id='userId']",
                "//main//input[contains(translate(@name,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'user')]",
                "//main//input[contains(translate(@placeholder,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'user')]",
                "(//main//input[@type='text' or @type='email' or not(@type)])[1]",
            ]
            for xpath in username_candidates:
                found = driver.find_elements(By.XPATH, xpath)
                username_input = next((elm for elm in found if elm.is_displayed()), None)
                if username_input:
                    break

            password_candidates = [
                "//main//input[@id='password']",
                "//main//input[contains(translate(@name,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'password')]",
                "(//main//input[@type='password'])[1]",
            ]
            for xpath in password_candidates:
                found = driver.find_elements(By.XPATH, xpath)
                password_input = next((elm for elm in found if elm.is_displayed()), None)
                if password_input:
                    break

            if not username_input or not password_input:
                return JsonResponse(
                    {
                        "status": "error",
                        "message": "Login form was not found on the ZabDesk portal.",
                    },
                    status=500,
                )

            # send_keys can be unreliable in attached debugger sessions; set values via JS.
            driver.execute_script(
                """
                arguments[0].value = arguments[1];
                arguments[0].dispatchEvent(new Event('input', { bubbles: true }));
                arguments[0].dispatchEvent(new Event('change', { bubbles: true }));
                """,
                username_input,
                "FAC001",
            )
            driver.execute_script(
                """
                arguments[0].value = arguments[1];
                arguments[0].dispatchEvent(new Event('input', { bubbles: true }));
                arguments[0].dispatchEvent(new Event('change', { bubbles: true }));
                """,
                password_input,
                "password123",
            )

            submitted_via_handler = driver.execute_script(
                """
                if (typeof window.handleLogin === 'function') {
                    window.handleLogin({ preventDefault: function () {} });
                    return true;
                }
                return false;
                """
            )

            if not submitted_via_handler:
                login_clicked = click_when_clickable(
                    "//main//button[contains(translate(normalize-space(.),'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'login')] | //main//input[@type='submit']",
                    "Login button",
                    timeout=8,
                )
                if not login_clicked:
                    password_input.send_keys(Keys.RETURN)

            try:
                wait_for_url(["roles.html", "faculty-courses.html", "course-hub.html", "recap-sheet.html"], timeout=12)
            except TimeoutException:
                return JsonResponse(
                    {
                        "status": "error",
                        "message": "Credentials were submitted, but portal login did not complete.",
                    },
                    status=500,
                )

        # Roles -> Faculty
        if "roles.html" in driver.current_url.lower():
            clicked_faculty = click_when_clickable(
                "//article[contains(@class, 'faculty-panel')] | //*[@role='button' and contains(., 'Faculty')] | //button[contains(., 'Faculty')]",
                "Faculty",
                timeout=8,
            )
            if not clicked_faculty:
                return JsonResponse(
                    {
                        "status": "error",
                        "message": "Logged in, but the Faculty button was not found.",
                    },
                    status=500,
                )
            wait_for_url(["faculty-courses.html"], timeout=10)

        # Faculty courses -> FOP
        if "faculty-courses.html" in driver.current_url.lower():
            clicked_fop = click_when_clickable(
                "//button[contains(normalize-space(.), 'FOP (CS101)')]",
                "FOP (CS101)",
                timeout=8,
            )
            if not clicked_fop:
                return JsonResponse(
                    {
                        "status": "error",
                        "message": "Logged in, but the FOP (CS101) course button was not found.",
                    },
                    status=500,
                )
            wait_for_url(["course-hub.html"], timeout=10)

        # Course hub -> Recap Sheet
        if "course-hub.html" in driver.current_url.lower():
            clicked_recap = click_when_clickable(
                "//button[contains(., 'Recap Sheet (Enter Marks)')]",
                "Recap Sheet (Enter Marks)",
                timeout=8,
            )
            if not clicked_recap:
                return JsonResponse(
                    {
                        "status": "error",
                        "message": "Course page loaded, but Recap Sheet (Enter Marks) was not found.",
                    },
                    status=500,
                )
            wait_for_url(["recap-sheet.html"], timeout=10)

        final_url = driver.current_url
        if "recap-sheet.html" not in final_url.lower():
            return JsonResponse(
                {
                    "status": "error",
                    "message": "Automation finished but recap sheet was not reached.",
                    "current_url": final_url,
                },
                status=500,
            )

        return JsonResponse(
            {
                "status": "success",
                "message": "ZabDesk login completed and recap sheet opened.",
                "current_url": final_url,
            }
        )
    except Exception as e:
        print(f"Zabdesk Auto Error: {e}")
        return JsonResponse({"status": "error", "message": str(e)}, status=500)


@csrf_exempt
def zabdesk_upload_excel(request):
    if request.method != "POST":
        return JsonResponse({"error": "Method not allowed"}, status=405)

    uploaded = request.FILES.get("file")
    if not uploaded:
        return JsonResponse({"error": "No file provided"}, status=400)

    allowed_ext = {".xlsx", ".xls", ".csv"}
    _, ext = os.path.splitext(uploaded.name or "")
    ext = ext.lower()
    if ext not in allowed_ext:
        return JsonResponse(
            {
                "error": "Unsupported file type. Please upload an Excel or CSV file.",
                "allowed": [".xlsx", ".xls", ".csv"],
            },
            status=400,
        )

    target_dir = os.path.join(tempfile.gettempdir(), "zabdesk_excel_uploads")
    os.makedirs(target_dir, exist_ok=True)

    temp_name = f"zabdesk_marks_{uuid.uuid4().hex}{ext}"
    temp_path = os.path.join(target_dir, temp_name)

    try:
        with open(temp_path, "wb") as tmp:
            for chunk in uploaded.chunks():
                tmp.write(chunk)

        old_path = request.session.get("zabdesk_uploaded_excel_path")
        if old_path and os.path.exists(old_path):
            try:
                os.remove(old_path)
            except Exception:
                pass

        request.session["zabdesk_uploaded_excel_path"] = temp_path
        request.session["zabdesk_uploaded_excel_name"] = uploaded.name
        request.session.pop("zabdesk_extracted_students", None)
        request.session.pop("zabdesk_extracted_metadata", None)
        request.session.modified = True

        return JsonResponse(
            {
                "status": "uploaded",
                "file_name": uploaded.name,
                "message": "Excel sheet uploaded successfully.",
            }
        )
    except Exception as e:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass
        return JsonResponse({"error": str(e)}, status=500)


@csrf_exempt
def zabdesk_extract_excel_data(request):
    if request.method not in ["GET", "POST"]:
        return JsonResponse({"error": "Method not allowed"}, status=405)

    file_path = request.session.get("zabdesk_uploaded_excel_path")
    file_name = request.session.get("zabdesk_uploaded_excel_name")

    if not file_path:
        return JsonResponse({"error": "No Excel sheet uploaded yet. Please upload a file first."}, status=400)

    if not os.path.exists(file_path):
        return JsonResponse({"error": "Uploaded file no longer exists. Please upload again."}, status=400)

    def _json_safe(value):
        if value is None:
            return None
        if isinstance(value, (datetime,)):
            return value.isoformat()
        if isinstance(value, (int, float, bool, str)):
            return value
        return str(value)

    _, ext = os.path.splitext(file_path)
    ext = ext.lower()

    try:
        structured_cache = _extract_zabdesk_students_with_marks(file_path)
        request.session["zabdesk_extracted_students"] = structured_cache.get("students", [])
        request.session["zabdesk_extracted_metadata"] = {
            "header_row": structured_cache.get("header_row"),
            "registration_column": structured_cache.get("registration_column"),
            "name_column": structured_cache.get("name_column"),
            "mark_columns": structured_cache.get("mark_columns", {}),
        }
        request.session.modified = True
    except Exception:
        request.session.pop("zabdesk_extracted_students", None)
        request.session.pop("zabdesk_extracted_metadata", None)
        request.session.modified = True

    try:
        if ext == ".csv":
            rows = []
            with open(file_path, "r", encoding="utf-8-sig", newline="") as f:
                reader = csv.reader(f)
                for row_idx, row in enumerate(reader, start=1):
                    cells = []
                    for col_idx, cell_val in enumerate(row, start=1):
                        cells.append(
                            {
                                "row": row_idx,
                                "column": col_idx,
                                "value": _json_safe(cell_val),
                            }
                        )
                    rows.append(cells)

            return JsonResponse(
                {
                    "status": "success",
                    "file_name": file_name,
                    "file_type": ext,
                    "sheet_count": 1,
                    "sheets": [
                        {
                            "name": "CSV",
                            "row_count": len(rows),
                            "rows": rows,
                        }
                    ],
                }
            )

        if ext in {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"}:
            try:
                from openpyxl import load_workbook
            except Exception as import_error:
                return JsonResponse(
                    {
                        "error": "Excel parser dependency is missing.",
                        "detail": str(import_error),
                    },
                    status=500,
                )

            workbook = load_workbook(filename=file_path, data_only=False)
            sheets_payload = []

            for sheet in workbook.worksheets:
                rows_payload = []
                for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
                    cell_row = []
                    for cell in row:
                        is_formula = isinstance(cell.value, str) and str(cell.value).startswith("=")
                        cell_row.append(
                            {
                                "coordinate": cell.coordinate,
                                "row": cell.row,
                                "column": cell.column,
                                "value": _json_safe(cell.value),
                                "data_type": cell.data_type,
                                "number_format": cell.number_format,
                                "is_formula": is_formula,
                            }
                        )
                    rows_payload.append(cell_row)

                sheets_payload.append(
                    {
                        "name": sheet.title,
                        "max_row": sheet.max_row,
                        "max_column": sheet.max_column,
                        "rows": rows_payload,
                    }
                )

            return JsonResponse(
                {
                    "status": "success",
                    "file_name": file_name,
                    "file_type": ext,
                    "sheet_count": len(sheets_payload),
                    "sheets": sheets_payload,
                }
            )

        return JsonResponse({"error": f"Unsupported file format for extraction: {ext}"}, status=400)
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)


ZABDESK_MARKS_TYPE_CONFIG = {
    "quiz_1": {
        "label": "Quiz 1 (4.5 marks)",
        "portal_option": "Quiz 1 (4.5 marks)",
        "portal_value": "quiz1",
        "header_aliases": ["quiz1", "quiz01", "q1", "quiz145marks"],
    },
    "quiz_2": {
        "label": "Quiz 2 (4.5 marks)",
        "portal_option": "Quiz 2 (4.5 marks)",
        "portal_value": "quiz2",
        "header_aliases": ["quiz2", "quiz02", "q2", "quiz245marks"],
    },
    "quiz_3": {
        "label": "Quiz 3 (4.5 marks)",
        "portal_option": "Quiz 3 (4.5 marks)",
        "portal_value": "quiz3",
        "header_aliases": ["quiz3", "quiz03", "q3", "quiz345marks"],
    },
    "quiz_4": {
        "label": "Quiz 4 (4.5 marks)",
        "portal_option": "Quiz 4 (4.5 marks)",
        "portal_value": "quiz4",
        "header_aliases": ["quiz4", "quiz04", "q4", "quiz445marks"],
    },
    "assignment_1": {
        "label": "Assignment 1 (3 marks)",
        "portal_option": "Assignment 1 (3 marks)",
        "portal_value": "assignment1",
        "header_aliases": ["assignment1", "assign1", "a1", "assignment13marks"],
    },
    "assignment_2": {
        "label": "Assignment 2 (3 marks)",
        "portal_option": "Assignment 2 (3 marks)",
        "portal_value": "assignment2",
        "header_aliases": ["assignment2", "assign2", "a2", "assignment23marks"],
    },
    "assignment_3": {
        "label": "Assignment 3 (3 marks)",
        "portal_option": "Assignment 3 (3 marks)",
        "portal_value": "assignment3",
        "header_aliases": ["assignment3", "assign3", "a3", "assignment33marks"],
    },
    "assignment_4": {
        "label": "Assignment 4 (3 marks)",
        "portal_option": "Assignment 4 (3 marks)",
        "portal_value": "assignment4",
        "header_aliases": ["assignment4", "assign4", "a4", "assignment43marks"],
    },
    "mid_term_paper": {
        "label": "Mid Term (30 marks)",
        "portal_option": "Mid Term (30 marks)",
        "portal_value": "midterm",
        "header_aliases": ["midterm", "midtermpaper", "midtermexam", "mid", "midterm30marks"],
    },
    "final_paper": {
        "label": "Final Term (40 marks)",
        "portal_option": "Final Term (40 marks)",
        "portal_value": "finalterm",
        "header_aliases": ["finalterm", "finalpaper", "finalexam", "final", "finalterm40marks"],
    },
}


def _normalize_name_for_match(value):
    text = str(value or "").strip().upper()
    text = re.sub(r"[^A-Z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _normalize_reg_for_match(value):
    return re.sub(r"[^A-Z0-9]+", "", str(value or "").strip().upper())


def _coerce_mark_value(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else round(value, 2)

    text = str(value).strip()
    if not text:
        return None

    normalized = text.replace(",", "")
    try:
        numeric = float(normalized)
        return int(numeric) if numeric.is_integer() else round(numeric, 2)
    except ValueError:
        return text


def _find_mark_column_index(normalized_headers, marks_type):
    config = ZABDESK_MARKS_TYPE_CONFIG.get(marks_type, {})
    aliases = list(config.get("header_aliases", []))
    aliases.append(config.get("label", ""))
    aliases.append(config.get("portal_option", ""))

    normalized_aliases = []
    for alias in aliases:
        normalized_alias = _normalize_header_text(alias)
        if normalized_alias:
            normalized_aliases.append(normalized_alias)

    for col_index, header_key in enumerate(normalized_headers):
        if not header_key:
            continue
        for alias in normalized_aliases:
            if alias in header_key or header_key in alias:
                return col_index

    # Semantic fallback for sheets where headers are split over multiple rows
    # (e.g., "Quiz" in one row and "1" in another row).
    mark_number_match = re.search(r"_(\d+)$", marks_type)
    mark_number = int(mark_number_match.group(1)) if mark_number_match else None

    for col_index, header_key in enumerate(normalized_headers):
        if not header_key:
            continue

        digits = re.findall(r"\d+", header_key)
        first_num = int(digits[0]) if digits else None

        if marks_type.startswith("quiz_"):
            if "quiz" in header_key and first_num == mark_number:
                return col_index
            if f"q{mark_number}" in header_key:
                return col_index

        if marks_type.startswith("assignment_"):
            if "assignment" in header_key and first_num == mark_number:
                return col_index
            if f"assign{mark_number}" in header_key or f"a{mark_number}" in header_key:
                return col_index

        if marks_type == "mid_term_paper":
            if "mid" in header_key and ("term" in header_key or "paper" in header_key):
                return col_index

        if marks_type == "final_paper":
            if "final" in header_key and ("term" in header_key or "paper" in header_key):
                return col_index
    return None


def _load_zabdesk_rows(file_path):
    _, ext = os.path.splitext(file_path)
    ext = ext.lower()

    if ext == ".csv":
        rows = []
        with open(file_path, "r", encoding="utf-8-sig", newline="") as csv_file:
            reader = csv.reader(csv_file)
            for row_number, row_values in enumerate(reader, start=1):
                rows.append({"row_number": row_number, "values": list(row_values)})
        return rows

    if ext in {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"}:
        try:
            from openpyxl import load_workbook
        except Exception as import_error:
            raise RuntimeError(f"Excel parser dependency is missing: {import_error}") from import_error

        workbook = load_workbook(filename=file_path, data_only=True, read_only=True)
        try:
            # Read every worksheet; header row may not be on the first tab.
            sheet_candidates = []
            for worksheet in workbook.worksheets:
                rows = []
                for row_number, row_values in enumerate(worksheet.iter_rows(values_only=True), start=1):
                    rows.append({"row_number": row_number, "values": list(row_values)})
                sheet_candidates.append({"sheet_name": worksheet.title, "rows": rows})
            return sheet_candidates
        finally:
            workbook.close()

    raise ValueError(f"Unsupported file format for extraction: {ext}")


def _extract_zabdesk_students_with_marks(file_path):
    loaded = _load_zabdesk_rows(file_path)

    # CSV returns plain rows; Excel returns per-sheet candidates.
    if loaded and isinstance(loaded, list) and loaded and isinstance(loaded[0], dict) and "sheet_name" in loaded[0]:
        sheet_candidates = loaded
    else:
        sheet_candidates = [{"sheet_name": "CSV", "rows": loaded or []}]

    best_match = None

    def _value_has_text(value):
        if value is None:
            return False
        return bool(str(value).strip())

    def _enrich_header_values(rows, header_index, base_values):
        # Some sheets split headers across multiple rows. Build a composite
        # header per column using nearby rows so patterns like "Quiz" + "1"
        # become matchable as "Quiz 1".
        enriched = list(base_values or [])
        max_cols = len(enriched)
        window_start = max(0, header_index - 4)
        window_end = min(len(rows), header_index + 2)

        for idx in range(window_start, window_end):
            row_values = rows[idx].get("values", [])
            if len(row_values) > max_cols:
                max_cols = len(row_values)

        if len(enriched) < max_cols:
            enriched.extend([None] * (max_cols - len(enriched)))

        composite = []
        for col_index in range(max_cols):
            tokens = []

            for lookback in range(window_start, window_end):
                row_values = rows[lookback].get("values", [])
                value = row_values[col_index] if col_index < len(row_values) else None
                if not _value_has_text(value):
                    continue

                token = str(value).strip()
                if token and token not in tokens:
                    tokens.append(token)

            if tokens:
                composite.append(" ".join(tokens))
            else:
                composite.append(enriched[col_index])

        return composite

    for candidate_sheet in sheet_candidates:
        rows = candidate_sheet.get("rows") or []
        if not rows:
            continue

        header_row_number = None
        registration_col = None
        name_col = None
        header_values = None
        normalized_headers = []

        for candidate_index, candidate in enumerate(rows[:120]):
            current_values = candidate.get("values", [])
            current_headers = [_normalize_header_text(value) for value in current_values]
            reg_idx = next((i for i, key in enumerate(current_headers) if _is_registration_header(key)), None)
            name_idx = next((i for i, key in enumerate(current_headers) if _is_student_name_header(key)), None)

            if reg_idx is not None and name_idx is not None:
                enriched_header_values = _enrich_header_values(rows, candidate_index, current_values)
                header_row_number = candidate.get("row_number")
                registration_col = reg_idx
                name_col = name_idx
                header_values = enriched_header_values
                normalized_headers = [_normalize_header_text(value) for value in enriched_header_values]
                break

        if header_row_number is None:
            continue

        mark_columns = {}
        matched_mark_count = 0
        for marks_type in ZABDESK_MARKS_TYPE_CONFIG:
            col_index = _find_mark_column_index(normalized_headers, marks_type)
            mark_columns[marks_type] = col_index
            if col_index is not None:
                matched_mark_count += 1

        sheet_score = {
            "matched_mark_count": matched_mark_count,
            "header_row": header_row_number,
        }

        if best_match is None or (
            sheet_score["matched_mark_count"], -sheet_score["header_row"]
        ) > (
            best_match["score"]["matched_mark_count"], -best_match["score"]["header_row"]
        ):
            best_match = {
                "sheet_name": candidate_sheet.get("sheet_name"),
                "rows": rows,
                "header_row": header_row_number,
                "registration_col": registration_col,
                "name_col": name_col,
                "header_values": header_values,
                "normalized_headers": normalized_headers,
                "mark_columns": mark_columns,
                "score": sheet_score,
            }

    if best_match is None:
        raise ValueError(
            "Unable to detect Student Name and Registration Number columns in uploaded sheet. "
            "Please ensure headers are present."
        )

    rows = best_match["rows"]
    header_row_number = best_match["header_row"]
    registration_col = best_match["registration_col"]
    name_col = best_match["name_col"]
    header_values = best_match["header_values"]
    mark_columns = best_match["mark_columns"]

    students = []
    for row_obj in rows:
        row_number = row_obj.get("row_number") or 0
        if row_number <= header_row_number:
            continue

        values = row_obj.get("values", [])
        registration_raw = values[registration_col] if registration_col < len(values) else None
        name_raw = values[name_col] if name_col < len(values) else None

        registration_no = str(registration_raw).strip() if registration_raw is not None else ""
        student_name = str(name_raw).strip() if name_raw is not None else ""
        if not registration_no and not student_name:
            continue

        marks_payload = {}
        for marks_type, col_index in mark_columns.items():
            if col_index is None or col_index >= len(values):
                marks_payload[marks_type] = None
                continue
            marks_payload[marks_type] = _coerce_mark_value(values[col_index])

        students.append(
            {
                "row_number": row_number,
                "registration_no": registration_no,
                "name": student_name,
                "marks": marks_payload,
            }
        )

    return {
        "sheet_name": best_match.get("sheet_name"),
        "header_row": header_row_number,
        "registration_column": registration_col + 1,
        "name_column": name_col + 1,
        "mark_columns": {
            marks_type: (column_index + 1) if column_index is not None else None
            for marks_type, column_index in mark_columns.items()
        },
        "headers": [str(value).strip() if value is not None else "" for value in (header_values or [])],
        "students": students,
    }


@csrf_exempt
def zabdesk_enter_marks_in_recap_sheet(request):
    if request.method != "POST":
        return JsonResponse({"error": "Method not allowed"}, status=405)

    run_user = _get_session_user(request)
    marks_type = ""
    marks_type_label = ""
    requested_count = 0
    success_count = 0
    failures = []
    source_file_name = request.session.get("zabdesk_uploaded_excel_name") or ""

    try:
        try:
            payload = json.loads(request.body.decode("utf-8") or "{}")
        except json.JSONDecodeError:
            return JsonResponse({"error": "Invalid JSON payload."}, status=400)

        marks_type = str((payload or {}).get("marks_type") or "").strip().lower()
        marks_config = ZABDESK_MARKS_TYPE_CONFIG.get(marks_type)
        if not marks_config:
            return JsonResponse(
                {
                    "error": "Invalid marks type selected.",
                    "allowed": [
                        {"value": key, "label": config["label"]}
                        for key, config in ZABDESK_MARKS_TYPE_CONFIG.items()
                    ],
                },
                status=400,
            )

        marks_type_label = marks_config["label"]

        file_path = request.session.get("zabdesk_uploaded_excel_path")
        if not file_path:
            return JsonResponse({"error": "No Excel sheet uploaded yet. Please upload a file first."}, status=400)

        if not os.path.exists(file_path):
            return JsonResponse({"error": "Uploaded file no longer exists. Please upload again."}, status=400)

        cached_students = request.session.get("zabdesk_extracted_students")
        cached_metadata = request.session.get("zabdesk_extracted_metadata") or {}
        structured_data = None

        if not isinstance(cached_students, list) or not cached_students:
            structured_data = _extract_zabdesk_students_with_marks(file_path)
            cached_students = structured_data.get("students", [])
            cached_metadata = {
                "header_row": structured_data.get("header_row"),
                "registration_column": structured_data.get("registration_column"),
                "name_column": structured_data.get("name_column"),
                "mark_columns": structured_data.get("mark_columns", {}),
            }
            request.session["zabdesk_extracted_students"] = cached_students
            request.session["zabdesk_extracted_metadata"] = cached_metadata
            request.session.modified = True

        mark_columns = (cached_metadata or {}).get("mark_columns") or {}
        if mark_columns.get(marks_type) is None:
            structured_data = _extract_zabdesk_students_with_marks(file_path)
            cached_students = structured_data.get("students", [])
            cached_metadata = {
                "header_row": structured_data.get("header_row"),
                "registration_column": structured_data.get("registration_column"),
                "name_column": structured_data.get("name_column"),
                "mark_columns": structured_data.get("mark_columns", {}),
            }
            request.session["zabdesk_extracted_students"] = cached_students
            request.session["zabdesk_extracted_metadata"] = cached_metadata
            request.session.modified = True

            mark_columns = (cached_metadata or {}).get("mark_columns") or {}
            if mark_columns.get(marks_type) is None:
                detected_headers = structured_data.get("headers") or []
                return JsonResponse(
                    {
                        "error": (
                            f"The selected marks column ({marks_type_label}) was not found in the uploaded sheet. "
                            "Please verify your header names and extract again."
                        ),
                        "detected_headers": detected_headers,
                        "expected_aliases": ZABDESK_MARKS_TYPE_CONFIG.get(marks_type, {}).get("header_aliases", []),
                    },
                    status=400,
                )

        selected_students = []
        for student in cached_students:
            if not isinstance(student, dict):
                continue
            registration_no = str(student.get("registration_no") or "").strip()
            student_name = str(student.get("name") or "").strip()
            if not registration_no and not student_name:
                continue

            marks_payload = student.get("marks") or {}
            selected_students.append(
                {
                    "registration_no": registration_no,
                    "name": student_name,
                    "mark": marks_payload.get(marks_type),
                    "row_number": student.get("row_number"),
                }
            )

        if not selected_students:
            return JsonResponse(
                {"error": "No student records were found in extracted data. Please extract the sheet again."},
                status=400,
            )

        requested_count = len(selected_students)

        options = Options()
        options.add_experimental_option("debuggerAddress", "127.0.0.1:9222")
        def _require_active_debug_browser():
            if _is_remote_debugger_running():
                return
            raise RuntimeError(
                "Zabdesk browser session is not active. Please click Login To Zabdesk first, then retry Enter Marks."
            )

        def _attach_driver_with_retry(max_attempts=2):
            last_error = None
            for _ in range(max_attempts):
                _require_active_debug_browser()
                tmp_driver = None
                try:
                    service = ChromeService(executable_path=get_driver_path())
                    tmp_driver = webdriver.Chrome(service=service, options=options)

                    handles = tmp_driver.window_handles
                    if not handles:
                        tmp_driver.execute_script("window.open('about:blank','_blank');")
                        handles = tmp_driver.window_handles

                    target_handle = handles[0]
                    for handle in handles:
                        tmp_driver.switch_to.window(handle)
                        current = (tmp_driver.current_url or "").lower()
                        if "hamdan-a11y.github.io/zabdesk-dummy-portal" in current:
                            target_handle = handle
                            break

                    tmp_driver.switch_to.window(target_handle)
                    return tmp_driver
                except (InvalidSessionIdException, WebDriverException) as attach_error:
                    last_error = attach_error
                    message = str(attach_error).lower()
                    if tmp_driver is not None:
                        try:
                            tmp_driver.quit()
                        except Exception:
                            pass

                    if "not connected to devtools" in message or "invalid session id" in message:
                        time.sleep(0.35)
                        continue
                    raise

            raise RuntimeError(
                "Chrome session disconnected. Please click Login To Zabdesk again and keep the browser window open."
            ) from last_error

        driver = _attach_driver_with_retry()

        recap_url = "https://hamdan-a11y.github.io/ZABDESK-Dummy-Portal/recap-sheet.html?courseId=1"
        driver.get(recap_url)
        time.sleep(1.1)

        current_url = (driver.current_url or "").lower()
        page_title = (driver.title or "").lower()
        if "index.html" in current_url or "login" in page_title:
            return JsonResponse(
                {
                    "error": "Portal session is not active. Please click Login To Zabdesk first.",
                    "code": "session_inactive",
                },
                status=400,
            )

        marks_dropdown = WebDriverWait(driver, 15).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "select"))
        )
        dropdown = Select(marks_dropdown)

        selected_option_text = ""
        portal_option_text = marks_config.get("portal_option", "")
        portal_option_value = str(marks_config.get("portal_value") or "").strip().lower()
        target_option_key = _normalize_header_text(portal_option_text)

        if portal_option_value:
            try:
                dropdown.select_by_value(portal_option_value)
                selected_option_text = (dropdown.first_selected_option.text or "").strip()
            except Exception:
                selected_option_text = ""

        for option in dropdown.options:
            if selected_option_text:
                break
            option_text = (option.text or "").strip()
            option_key = _normalize_header_text(option_text)
            if portal_option_text.lower() in option_text.lower() or (
                target_option_key and (target_option_key in option_key or option_key in target_option_key)
            ):
                selected_option_text = option_text
                dropdown.select_by_visible_text(option_text)
                break

        if not selected_option_text:
            return JsonResponse(
                {
                    "error": (
                        f"Could not find the selected marks type ({marks_type_label}) "
                        "on the ZabDesk recap sheet dropdown."
                    )
                },
                status=500,
            )

        try:
            driver.execute_script(
                "arguments[0].dispatchEvent(new Event('input', { bubbles: true }));"
                "arguments[0].dispatchEvent(new Event('change', { bubbles: true }));",
                marks_dropdown,
            )
        except Exception:
            pass

        time.sleep(0.4)
        open_result = driver.execute_script(
            "if (typeof openEnterMarks === 'function') { openEnterMarks(); return 'fn'; }"
            "var btn = Array.from(document.querySelectorAll('button')).find(function(b){"
            "  return /enter\\s*marks/i.test((b.textContent || '').trim());"
            "});"
            "if (btn) { btn.click(); return 'btn'; }"
            "return 'none';"
        )
        if open_result == "none":
            return JsonResponse(
                {"error": "Could not open marks entry table on the recap sheet."},
                status=500,
            )
        time.sleep(0.8)

        WebDriverWait(driver, 20).until(EC.presence_of_element_located((By.CSS_SELECTOR, "table tbody tr")))

        portal_rows = driver.find_elements(By.CSS_SELECTOR, "table tbody tr")
        portal_by_reg = {}
        portal_by_name = {}
        portal_by_reg_name = {}

        for row in portal_rows:
            cells = row.find_elements(By.CSS_SELECTOR, "td")
            if len(cells) < 4:
                continue

            registration_text = (cells[0].text or "").strip()
            name_text = (cells[1].text or "").strip()

            try:
                input_element = cells[-1].find_element(By.CSS_SELECTOR, "input")
            except NoSuchElementException:
                continue

            reg_key = _normalize_reg_for_match(registration_text)
            name_key = _normalize_name_for_match(name_text)
            candidate = {
                "registration_no": registration_text,
                "name": name_text,
                "input": input_element,
            }

            if reg_key:
                portal_by_reg.setdefault(reg_key, []).append(candidate)
            if name_key:
                portal_by_name.setdefault(name_key, []).append(candidate)
            if reg_key and name_key:
                portal_by_reg_name[f"{reg_key}|{name_key}"] = candidate

        def resolve_portal_student(registration_no, student_name):
            reg_key = _normalize_reg_for_match(registration_no)
            name_key = _normalize_name_for_match(student_name)

            if reg_key and name_key:
                direct = portal_by_reg_name.get(f"{reg_key}|{name_key}")
                if direct:
                    return direct

            if reg_key:
                by_reg = portal_by_reg.get(reg_key, [])
                if by_reg:
                    if name_key:
                        for candidate in by_reg:
                            if _normalize_name_for_match(candidate.get("name")) == name_key:
                                return candidate
                    return by_reg[0]

            if name_key:
                by_name = portal_by_name.get(name_key, [])
                if by_name:
                    return by_name[0]

            return None

        for student in selected_students:
            registration_no = student.get("registration_no") or ""
            student_name = student.get("name") or ""
            mark_value = student.get("mark")

            if mark_value in (None, ""):
                failures.append(
                    {
                        "registration_no": registration_no,
                        "name": student_name,
                        "reason": "Mark value is missing in the uploaded sheet.",
                    }
                )
                continue

            if isinstance(mark_value, str):
                stripped_mark = mark_value.strip()
                if not re.fullmatch(r"-?\d+(\.\d+)?", stripped_mark):
                    failures.append(
                        {
                            "registration_no": registration_no,
                            "name": student_name,
                            "reason": f"Mark value '{mark_value}' is not numeric.",
                        }
                    )
                    continue
                mark_text = stripped_mark
            else:
                mark_text = str(mark_value)

            target_student = resolve_portal_student(registration_no, student_name)
            if not target_student:
                failures.append(
                    {
                        "registration_no": registration_no,
                        "name": student_name,
                        "reason": "Student not found on recap sheet.",
                    }
                )
                continue

            try:
                input_element = target_student["input"]
                driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", input_element)
                time.sleep(0.08)
                input_element.click()
                time.sleep(0.05)
                input_element.send_keys(Keys.CONTROL, "a")
                input_element.send_keys(Keys.BACKSPACE)
                input_element.send_keys(mark_text)
                driver.execute_script(
                    "arguments[0].dispatchEvent(new Event('input', { bubbles: true }));"
                    "arguments[0].dispatchEvent(new Event('change', { bubbles: true }));",
                    input_element,
                )
                time.sleep(0.1)
                success_count += 1
            except Exception as fill_error:
                fill_reason = str(fill_error)
                fill_reason_lower = fill_reason.lower()
                if (
                    "stacktrace:" in fill_reason_lower
                    or "chromedriver" in fill_reason_lower
                    or "gethandleverifier" in fill_reason_lower
                    or "invalid session id" in fill_reason_lower
                ):
                    fill_reason = "Browser session disconnected during input."
                failures.append(
                    {
                        "registration_no": registration_no,
                        "name": student_name,
                        "reason": f"Failed to enter mark: {fill_reason}",
                    }
                )

        failure_count = len(failures)
        status_value = "completed" if failure_count == 0 else "completed_with_errors"

        try:
            ZabdeskMarkEntryRun.objects.create(
                user=run_user,
                marks_type=marks_type,
                marks_type_label=marks_type_label,
                source_file_name=source_file_name,
                requested_count=requested_count,
                success_count=success_count,
                failure_count=failure_count,
                failures=failures,
                status=status_value,
            )
        except Exception as db_error:
            print(f"Zabdesk marks run log save failed: {db_error}")

        if failure_count == 0:
            completion_message = (
                f"{marks_type_label} marks have been successfully entered for all {success_count} students."
            )
        else:
            completion_message = (
                f"{marks_type_label} marks were entered for {success_count} of {requested_count} students. "
                f"{failure_count} student records were skipped."
            )

        return JsonResponse(
            {
                "status": status_value,
                "selected_marks_type": marks_type_label,
                "requested_count": requested_count,
                "success_count": success_count,
                "failure_count": failure_count,
                "failures": failures,
                "message": completion_message,
            }
        )

    except Exception as e:
        error_text = str(e)
        lowered_error = error_text.lower()

        if (
            "invalid session id" in lowered_error
            or "not connected to devtools" in lowered_error
            or "chrome not reachable" in lowered_error
            or "target window already closed" in lowered_error
        ):
            error_text = "Chrome session disconnected. Please click Login To Zabdesk again and keep the browser window open."
        elif "chromedriver" in lowered_error or "stacktrace:" in lowered_error or "gethandleverifier" in lowered_error:
            error_text = "Browser automation failed while operating the recap sheet. Please click Login To Zabdesk again, then retry Enter Marks."

        try:
            ZabdeskMarkEntryRun.objects.create(
                user=run_user,
                marks_type=marks_type or "unknown",
                marks_type_label=marks_type_label or "Unknown",
                source_file_name=source_file_name,
                requested_count=requested_count,
                success_count=success_count,
                failure_count=len(failures),
                failures=failures,
                status="failed",
                error_message=error_text,
            )
        except Exception as db_error:
            print(f"Zabdesk marks failed-run log save failed: {db_error}")

        status_code = 400 if "chrome session disconnected" in error_text.lower() else 500
        return JsonResponse({"error": error_text}, status=status_code)

